"""
报表生成工具集 (Model 层 / 工具箱) — 对接清单 #15/#16/#17

四个工具:
  1. generate_change_stats_chart — 变化统计图表 (matplotlib 饼图+柱图)
  2. generate_monitor_report    — 监测分析报告 (Jinja2 模板 + reportlab PDF / python-docx Word)
  3. export_change_vector       — 变化图斑矢量导出 (GeoJSON → Shapefile ZIP)
  4. export_stats_table         — 统计表格导出 (Excel/CSV)

★ 重工具白名单 (HEAVY_TOOL_CATEGORIES 含 "report"), 调用会触发 ai_task 持久化.
★ 依赖策略: 全部走宿主机 (sam3 环境), 缺失时返回 error.
  - matplotlib/pandas: 通常已装 (samseg/sandbox 都依赖)
  - reportlab/python-docx/jinja2: 报告生成专用, 需 pip install
  - geopandas/shapely: 矢量导出用, 通常已装

依赖方向: model.tools.report_tools → data.vector (复用) / 宿主机 matplotlib+reportlab
"""
import os
import re
import time
import logging
from pathlib import Path
from typing import Dict, Any, Optional

from langchain_core.tools import tool
from backend.model.tools.tool_registry import register_tool
from backend.model.tools._result import build_render_image_action, build_download_action
from backend.model.tools._shapefile import ascii_stem
from backend.model.tools._verification import verify_by_path
from backend.model.tools._paths import session_subpath, short_conv, rand_suffix
from backend.config import settings

logger = logging.getLogger(__name__)


# ★ _ascii_safe 已统一为 backend.model.tools._shapefile.ascii_stem (复用 geoio 公共函数)


def _get_report_output_dir() -> Path:
    """★ v2.4: 报表产物输出目录: agent-files/report/{proj}/{conv}/"""
    d = Path(settings.file_storage_root).resolve() / "report" / session_subpath()
    d.mkdir(parents=True, exist_ok=True)
    return d


def _build_download_url(file_path: Path) -> str:
    """★ C4 归一: 复用 _paths.build_download_url 统一实现 (消除重复的 relative_to 逻辑)."""
    from backend.model.tools._paths import build_download_url as _unified
    return _unified(file_path, Path(settings.file_storage_root))


def _load_change_geojson(geojson_path: str) -> Optional["gpd.GeoDataFrame"]:
    """加载变化图斑 GeoJSON, 失败返回 None."""
    try:
        import geopandas as gpd
    except ImportError:
        logger.warning("[Report] geopandas 未安装")
        return None
    if not Path(geojson_path).is_file():
        logger.warning(f"[Report] GeoJSON 不存在: {geojson_path}")
        return None
    try:
        return gpd.read_file(geojson_path)
    except Exception as e:
        logger.warning(f"[Report] GeoJSON 读取失败: {e}")
        return None


def _summarize_gdf(gdf) -> Dict[str, Any]:
    """从 GeoDataFrame 提取分类统计 (供图表/报告用)."""
    per_class = {}
    for _, row in gdf.iterrows():
        name = row.get("class_name") or row.get("class") or "未知"
        if name not in per_class:
            per_class[name] = {"name": name, "count": 0, "area_m2": 0.0}
        per_class[name]["count"] += 1
        area = row.get("area_m2")
        if area is None:
            area = float(row.geometry.area) if row.geometry else 0.0
        per_class[name]["area_m2"] += float(area)
    per_class_list = sorted(per_class.values(), key=lambda x: -x["area_m2"])
    total_area = sum(p["area_m2"] for p in per_class_list)
    return {
        "per_class": per_class_list,
        "total_features": len(gdf),
        "total_area_m2": round(total_area, 2),
    }


# ==================== 1. 变化统计图表 ====================

@register_tool("report")
@tool
def generate_change_stats_chart(change_geojson_path: str, chart_type: str = "both") -> Dict[str, Any]:
    """
    基于变化图斑 GeoJSON 生成统计图表 (饼图+柱图), 用于可视化各类变化占比 (#15)。

    ★ 实现: matplotlib 绘制, 中文标签用系统 CJK 字体. 输出 PNG, 经 render_sandbox_image 展示。
    ★ 耗时: 通常 1~3 秒。

    入参:
        - change_geojson_path (str): 变化图斑 GeoJSON 路径 (detect_change 产出的矢量)
        - chart_type (str): 图表类型, "pie" (饼图) / "bar" (柱图) / "both" (并排两张, 默认)
    出参: {type: frontend_action, instruction: render_sandbox_image, summary, data, verification}

    ★ 成功契约 (全部满足才算完成, 否则视为失败):
      1. type == "frontend_action" (而非 "error")
      2. verification.ok == True (PNG 文件真实生成且 > 100B)
      3. data.stats.total_features > 0 (源图斑非空)
      任一不满足 → 必须按失败处理, 不可向用户汇报"已生成图表"
    示例:
        generate_change_stats_chart("E:/.../change_result.geojson")
        generate_change_stats_chart("E:/.../change_result.geojson", "pie")
    """
    gdf = _load_change_geojson(change_geojson_path)
    if gdf is None:
        return {"type": "error", "msg": f"无法读取变化图斑: {change_geojson_path}"}
    if gdf.empty:
        return {"type": "error", "msg": "变化图斑为空, 无法生成图表"}

    try:
        import matplotlib
        matplotlib.use("Agg")  # 非交互后端
        import matplotlib.pyplot as plt
        # ★ 中文字体: 优先用系统已装的 CJK 字体, 避免图表中文显示成方块
        import matplotlib.font_manager as fm
        _CJK_FONT_CANDIDATES = [
            "Microsoft YaHei", "SimHei", "SimSun", "KaiTi",
            "Noto Sans CJK SC", "Noto Sans CJK JP", "WenQuanYi Zen Hei",
            "Source Han Sans SC", "PingFang SC",
        ]
        # 找到系统真实存在的字体
        _available_fonts = {f.name for f in fm.fontManager.ttflist}
        _chosen = None
        for fname in _CJK_FONT_CANDIDATES:
            if fname in _available_fonts:
                _chosen = fname
                break
        if _chosen:
            matplotlib.rcParams["font.sans-serif"] = [_chosen] + matplotlib.rcParams.get("font.sans-serif", [])
            matplotlib.rcParams["axes.unicode_minus"] = False
        else:
            # 兜底: 直接加载 Windows 字体文件
            for fp in [r"C:\Windows\Fonts\msyh.ttc", r"C:\Windows\Fonts\simhei.ttf",
                       "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"]:
                if os.path.isfile(fp):
                    fm.fontManager.addfont(fp)
                    matplotlib.rcParams["font.sans-serif"] = [fm.FontProperties(fname=fp).get_name()] + matplotlib.rcParams.get("font.sans-serif", [])
                    matplotlib.rcParams["axes.unicode_minus"] = False
                    break
    except ImportError:
        return {"type": "error", "msg": "matplotlib 未安装, 无法生成图表"}

    stats = _summarize_gdf(gdf)
    per_class = stats["per_class"]
    names = [p["name"] for p in per_class]
    counts = [p["count"] for p in per_class]
    areas = [p["area_m2"] for p in per_class]

    out_dir = _get_report_output_dir()
    # ★ v2.4 命名精简: 不再吃上游长 stem, 改用 {conv短}_stats_{pie|bar}_{随机4}.png
    stem = (short_conv() or "anon") + "_stats"
    rs = rand_suffix(4)

    chart_files = []
    summary_parts = []

    # 饼图 (按面积占比)
    if chart_type in ("pie", "both"):
        pie_path = out_dir / f"{stem}_pie_{rs}.png"
        fig, ax = plt.subplots(figsize=(7, 6))
        total = sum(areas) or 1
        ax.pie(areas, labels=[f"{n}\n{a/total*100:.1f}%" for n, a in zip(names, areas)],
               autopct=None, startangle=90, textprops={"fontsize": 10})
        ax.set_title("变化面积占比", fontsize=14)
        fig.tight_layout()
        fig.savefig(str(pie_path), dpi=120, bbox_inches="tight")
        plt.close(fig)
        chart_files.append(("饼图", pie_path))
        summary_parts.append(f"饼图: 各类变化面积占比")

    # 柱图 (按图斑数+面积)
    if chart_type in ("bar", "both"):
        bar_path = out_dir / f"{stem}_bar_{rs}.png"
        fig, ax1 = plt.subplots(figsize=(9, 6))
        x = range(len(names))
        bars = ax1.bar(x, counts, color="#4a90e2", alpha=0.75, label="图斑数")
        ax1.set_xlabel("变化类别")
        ax1.set_ylabel("图斑数", color="#4a90e2")
        ax1.set_xticks(list(x))
        ax1.set_xticklabels(names, rotation=20, ha="right")
        ax2 = ax1.twinx()
        ax2.plot(x, areas, color="#e74c3c", marker="o", linewidth=2, label="面积 (m²)")
        ax2.set_ylabel("面积 (m²)", color="#e74c3c")
        ax1.set_title("变化类别统计 (图斑数 + 面积)", fontsize=14)
        fig.tight_layout()
        fig.savefig(str(bar_path), dpi=120, bbox_inches="tight")
        plt.close(fig)
        chart_files.append(("柱图", bar_path))
        summary_parts.append(f"柱图: 各类变化图斑数与面积对照")

    if not chart_files:
        return {"type": "error", "msg": f"未生成图表 (chart_type={chart_type} 无效)"}

    # 构造 frontend_action: 多图渲染 (沿用 render_image 协议)
    # 单图用 render_image, 多图用多个 image_url
    params = {}
    if len(chart_files) == 1:
        label, p = chart_files[0]
        url = _build_download_url(p) + f"?t={rs}"
        params = {"image_url": url, "caption": f"变化统计{label}: {stem}"}
    else:
        # 多图: 第一张走 image_url, 第二张走 image_url_extra
        urls = [{"url": _build_download_url(p) + f"?t={rs}", "caption": f"变化统计{label}: {stem}"}
                for label, p in chart_files]
        params = {
            "image_url": urls[0]["url"],
            "caption": urls[0]["caption"],
            "extra_images": urls[1:],
        }

    summary_text = (
        f"已生成变化统计图表 ({', '.join(summary_parts)})。\n"
        f"统计概要: 共 {stats['total_features']} 个图斑, 总面积 {stats['total_area_m2']} m²"
        f" ({stats['total_area_m2']/10000:.2f} 公顷)\n"
        f"各类详情:\n"
    )
    for p in per_class[:6]:
        summary_text += f"  - {p['name']}: {p['count']}个, {p['area_m2']} m²\n"

    # ★ 成功契约校验: 每张 PNG 都必须真实落地且非空, 否则视为失败
    verifications = []
    all_ok = True
    for label, p in chart_files:
        v = verify_by_path(p)
        v["chart"] = label
        verifications.append(v)
        if not v["ok"]:
            all_ok = False
    verification_summary = {
        "ok": all_ok,
        "charts": verifications,
        "evidence": (f"{len(verifications)} 张图表全部通过校验"
                     if all_ok else f"{sum(1 for v in verifications if not v['ok'])} 张图表校验失败"),
    }

    if not all_ok:
        # ★ 不准伪装成功: 校验失败直接报错, 让 AI 走反思流程
        failed = [v for v in verifications if not v["ok"]]
        return {
            "type": "error",
            "msg": f"图表生成失败: {failed[0].get('reason', '未知原因')}",
            "verification": verification_summary,
            "data": {"stats": stats, "chart_files": [str(p) for _, p in chart_files]},
        }

    return build_render_image_action(
        params=params,
        summary=summary_text,
        description=f"⏳ 统计图表已生成 ({len(chart_files)} 张)",
        verification=verification_summary,
        data={
            "stats": stats,
            "chart_files": [str(p) for _, p in chart_files],
        },
    )


# ==================== 2. 监测报告生成 ====================

@register_tool("report")
@tool
def generate_monitor_report(
    change_geojson_path: str,
    output_format: str = "pdf",
    region_name: str = "",
    t1_name: str = "",
    t2_name: str = "",
    change_type: str = "",
    change_confidence: float = 0.0,
    change_reason: str = "",
    overlay_image_path: str = "",
    overlay_image_caption: str = "",
) -> Dict[str, Any]:
    """
    基于变化图斑数据生成标准化监测分析报告, 支持 PDF / Word / Markdown 三种格式 (#16)。

    ★ 报告内容: 变化概况 + 各类别统计表 + 主要变化分析 + 建议与后续工作
    ★ 实现: Jinja2 渲染 Markdown → reportlab 转 PDF / python-docx 转 Word
    ★ 耗时: 通常 1~5 秒。

    入参:
        - change_geojson_path (str): 变化图斑 GeoJSON 路径
        - output_format (str): 输出格式, "pdf" (默认) / "word" / "markdown"
        - region_name (str): 监测区域名称 (报告抬头用)
        - t1_name / t2_name (str): T1/T2 影像名称
        - change_type (str): 业务变化类型 (来自 detect_change 的 change_type)
        - change_confidence (float): 业务判读置信度 (0~1)
        - change_reason (str): 业务判读依据
        - overlay_image_path (str): 可选, 原图叠加矢量图斑的 PNG 路径
        - overlay_image_caption (str): 可选, 叠加图说明文字
    出参: {type: frontend_action, instruction: download, summary, data, verification}

    ★ 成功契约 (全部满足才算完成, 否则视为失败):
      1. type == "frontend_action" (而非 "error")
      2. verification.ok == True (PDF/Word 文件可被读回且页数 > 0)
      3. data.stats.total_features > 0 (源图斑非空)
      任一不满足 → 必须按失败处理, 不可向用户汇报"报告已生成"
    示例:
        generate_monitor_report("E:/.../change.geojson", "pdf", "成都某区", change_type="新增建筑")
    """
    gdf = _load_change_geojson(change_geojson_path)
    if gdf is None:
        return {"type": "error", "msg": f"无法读取变化图斑: {change_geojson_path}"}

    stats = _summarize_gdf(gdf)
    per_class = stats["per_class"]
    dominant = per_class[0] if per_class else None
    overlay_path = Path(overlay_image_path).resolve() if overlay_image_path else None
    overlay_verification = None
    if overlay_path is not None:
        overlay_verification = verify_by_path(overlay_path)
        if not overlay_verification["ok"]:
            return {
                "type": "error",
                "msg": f"报告叠加图不可用: {overlay_verification.get('reason', '校验未通过')}",
                "verification": overlay_verification,
                "data": {
                    "overlay_image_path": str(overlay_path),
                    "format": output_format,
                    "stats": stats,
                },
            }

    # 准备 Jinja2 上下文
    context = {
        "generated_at": time.strftime("%Y-%m-%d %H:%M:%S"),
        "region_name": region_name,
        "t1_name": t1_name or Path(change_geojson_path).stem + "_t1",
        "t2_name": t2_name or Path(change_geojson_path).stem + "_t2",
        "change_type": change_type,
        "change_confidence": change_confidence,
        "change_reason": change_reason,
        "total_features": stats["total_features"],
        "total_area_m2": stats["total_area_m2"],
        "total_area_ha": round(stats["total_area_m2"] / 10000, 2),
        "changed_percent": "—",  # 无监测区总面积时留空
        "per_class": per_class,
        "dominant_class": dominant,
        "overlay_image_path": str(overlay_path) if overlay_path else "",
        "overlay_image_caption": overlay_image_caption,
    }

    # 渲染 Markdown
    try:
        from jinja2 import Environment, FileSystemLoader
    except ImportError:
        return {"type": "error", "msg": "jinja2 未安装, 无法渲染报告模板"}

    template_dir = Path(__file__).resolve().parent.parent / "report_templates"
    env = Environment(loader=FileSystemLoader(str(template_dir)), autoescape=False)
    template = env.get_template("monitor_report.md.j2")
    md_content = template.render(**context)

    out_dir = _get_report_output_dir()
    # ★ v2.4 命名精简: 不再吃上游长 stem (套娃根因), 改用 {conv短}_report_{随机4}
    #   ASCII 化仍保留 (PDF/docx 库在 Windows + 中文路径可能静默失败)
    safe_stem = ascii_stem(short_conv() or "anon", fallback="anon") + "_report"
    rs = rand_suffix(4)
    output_format = output_format.lower()

    if output_format == "markdown":
        out_path = out_dir / f"{safe_stem}_{rs}.md"
        out_path.write_text(md_content, encoding="utf-8")
    elif output_format == "pdf":
        try:
            out_path = out_dir / f"{safe_stem}_{rs}.pdf"
            _md_to_pdf(md_content, str(out_path))
        except Exception as e:
            logger.warning(f"[Report] PDF 生成失败, 降级 markdown: {e}")
            out_path = out_dir / f"{safe_stem}_{rs}.md"
            out_path.write_text(md_content, encoding="utf-8")
            output_format = "markdown (PDF 降级)"
    elif output_format == "word":
        try:
            out_path = out_dir / f"{safe_stem}_{rs}.docx"
            _md_to_docx(md_content, str(out_path))
        except Exception as e:
            logger.warning(f"[Report] Word 生成失败, 降级 markdown: {e}")
            out_path = out_dir / f"{safe_stem}_{rs}.md"
            out_path.write_text(md_content, encoding="utf-8")
            output_format = "markdown (Word 降级)"
    else:
        return {"type": "error", "msg": f"output_format 无效: {output_format} (pdf/word/markdown)"}

    download_url = _build_download_url(out_path) + f"?t={rs}"
    size = out_path.stat().st_size

    # ★ 成功契约校验: PDF/Word/markdown 必须可被读回且非空
    verification = verify_by_path(out_path)

    if not verification["ok"]:
        return {
            "type": "error",
            "msg": f"报告生成失败: {verification.get('reason', '校验未通过')}",
            "verification": verification,
            "data": {"output_path": str(out_path), "format": output_format, "stats": stats},
        }

    summary = (
        f"已生成监测分析报告: {out_path.name} ({output_format}, {size//1024}KB)\n"
        f"报告概要: {stats['total_features']} 个图斑, 总面积 {stats['total_area_m2']} m²"
        f" ({stats['total_area_m2']/10000:.2f} 公顷)\n"
        f"主导变化: {dominant['name'] if dominant else '无'}\n"
        f"叠加图: {'已嵌入' if overlay_path else '未提供'}\n"
        f"业务类型: {change_type or '未指定'}\n"
        f"下载链接已生成, 请在前端下载查看。"
    )

    return build_download_action(
        params={"url": download_url, "filename": out_path.name, "caption": f"监测报告 ({output_format})"},
        summary=summary,
        description=f"⏳ 报告已生成 ({output_format})",
        verification=verification,
        data={
            "artifact_path": str(out_path),
            "artifact_url": download_url,
            "format": output_format,
            "stats": stats,
            "overlay_image_path": str(overlay_path) if overlay_path else "",
            "overlay_image_verification": overlay_verification,
        },
    )


def _parse_markdown_image(line: str) -> Optional[Dict[str, str]]:
    """
    入参:
      - line: Markdown 单行文本。
    方法:
      - 解析形如 ![caption](path) 的图片语法。
      - 仅返回本地文件路径, 后续渲染器负责校验文件是否存在。
    出参:
      - dict, 包含 caption 与 path; 非图片行返回 None。
    """
    match = re.match(r"^!\[(?P<caption>[^\]]*)\]\((?P<path>[^)]+)\)\s*$", line.strip())
    if not match:
        return None
    return {
        "caption": match.group("caption").strip(),
        "path": match.group("path").strip().strip('"'),
    }


def _append_pdf_image(story: list, image_path: str, caption: str, body_style: Any, max_width: float) -> None:
    """
    入参:
      - story: reportlab flowable 列表。
      - image_path: 本地图片路径。
      - caption: 图片标题。
      - body_style: PDF 正文字体样式。
      - max_width: 图片最大显示宽度。
    方法:
      - 按页面宽度等比例缩放图片。
      - 图片缺失时写入一行证据提示, 不让 PDF 渲染器崩溃。
    出参:
      - None。
    """
    from reportlab.platypus import Image as RLImage, Paragraph, Spacer
    image_file = Path(image_path)
    if not image_file.is_file():
        story.append(Paragraph(f"图片不存在: {image_path}", body_style))
        return
    try:
        from PIL import Image as PILImage
        with PILImage.open(image_file) as img:
            width, height = img.size
        if width <= 0 or height <= 0:
            story.append(Paragraph(f"图片尺寸无效: {image_path}", body_style))
            return
        draw_width = min(max_width, float(width))
        draw_height = draw_width * float(height) / float(width)
        max_height = 300.0
        if draw_height > max_height:
            draw_height = max_height
            draw_width = draw_height * float(width) / float(height)
        story.append(RLImage(str(image_file), width=draw_width, height=draw_height))
        if caption:
            safe_caption = caption.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            story.append(Paragraph(safe_caption, body_style))
        story.append(Spacer(1, 8))
    except Exception as exc:
        story.append(Paragraph(f"图片嵌入失败: {exc}", body_style))


def _md_to_pdf(md_content: str, output_path: str) -> None:
    """
    入参:
      - md_content: 已渲染的 Markdown 文本。
      - output_path: PDF 输出路径。
    方法:
      - 使用 reportlab 将标题、表格、正文和本地图片转换为 PDF。
      - 图片语法支持 ![caption](local_path)。
    出参:
      - None。
    """
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib import colors
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    # 注册中文字体 (Windows: C:\Windows\Fonts\msyh.ttc / simhei.ttf; Linux: noto-cjk)
    font_name = "Helvetica"  # 默认 fallback
    font_paths = [
        ("CJK", r"C:\Windows\Fonts\msyh.ttc"),
        ("CJK", r"C:\Windows\Fonts\simhei.ttf"),
        ("CJK", r"C:\Windows\Fonts\simsun.ttc"),
        ("CJK", "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
        ("CJK", "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc"),
    ]
    for name, path in font_paths:
        try:
            if os.path.isfile(path):
                pdfmetrics.registerFont(TTFont(name, path))
                font_name = name
                break
        except Exception:
            continue

    doc = SimpleDocTemplate(output_path, pagesize=A4,
                            topMargin=50, bottomMargin=50, leftMargin=50, rightMargin=50)
    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("H1CJK", parent=styles["Heading1"], fontName=font_name, fontSize=18, spaceAfter=14)
    h2 = ParagraphStyle("H2CJK", parent=styles["Heading2"], fontName=font_name, fontSize=14, spaceAfter=10)
    body = ParagraphStyle("BodyCJK", parent=styles["BodyText"], fontName=font_name, fontSize=10.5, leading=18)

    story = []
    lines = md_content.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if line.startswith("# "):
            story.append(Paragraph(line[2:].strip(), h1))
            story.append(Spacer(1, 8))
        elif line.startswith("## "):
            story.append(Paragraph(line[3:].strip(), h2))
            story.append(Spacer(1, 6))
        elif line.startswith("|") and i + 1 < len(lines) and lines[i+1].startswith("|"):
            # 简单表格解析
            table_rows = []
            while i < len(lines) and lines[i].startswith("|"):
                row = [c.strip() for c in lines[i].strip("|").split("|")]
                if not all(set(c) <= set("-: ") for c in row):  # 跳过分隔行
                    table_rows.append(row)
                i += 1
            if table_rows:
                t = Table(table_rows, hAlign="LEFT")
                t.setStyle(TableStyle([
                    ("FONTNAME", (0, 0), (-1, -1), font_name),
                    ("FONTSIZE", (0, 0), (-1, -1), 9),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4a90e2")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f0f7ff")]),
                ]))
                story.append(t)
                story.append(Spacer(1, 8))
            continue
        elif _parse_markdown_image(line):
            image_info = _parse_markdown_image(line) or {}
            _append_pdf_image(
                story,
                image_info.get("path", ""),
                image_info.get("caption", ""),
                body,
                doc.width,
            )
        elif line.startswith("---"):
            story.append(Spacer(1, 8))
        elif line.strip():
            # 转义 XML 特殊字符
            text = line.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            story.append(Paragraph(text, body))
        else:
            story.append(Spacer(1, 6))
        i += 1

    doc.build(story)


def _md_to_docx(md_content: str, output_path: str) -> None:
    """
    入参:
      - md_content: 已渲染的 Markdown 文本。
      - output_path: Word 输出路径。
    方法:
      - 使用 python-docx 将标题、表格、正文和本地图片转换为 docx。
      - 图片语法支持 ![caption](local_path)。
    出参:
      - None。
    """
    from docx import Document
    from docx.shared import Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    lines = md_content.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if line.startswith("# "):
            p = doc.add_heading(line[2:].strip(), level=1)
        elif line.startswith("## "):
            p = doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("|") and i + 1 < len(lines) and lines[i+1].startswith("|"):
            # 简单表格
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip("|").split("|")]
                if not all(set(c) <= set("-: ") for c in cells):
                    rows.append(cells)
                i += 1
            if rows:
                table = doc.add_table(rows=len(rows), cols=len(rows[0]))
                table.style = "Light Grid Accent 1"
                for r_idx, row in enumerate(rows):
                    for c_idx, cell_text in enumerate(row):
                        if c_idx < len(table.rows[r_idx].cells):
                            table.rows[r_idx].cells[c_idx].text = cell_text
            continue
        elif _parse_markdown_image(line):
            image_info = _parse_markdown_image(line) or {}
            image_path = image_info.get("path", "")
            if Path(image_path).is_file():
                doc.add_picture(image_path, width=Inches(6))
                if image_info.get("caption"):
                    p = doc.add_paragraph(image_info["caption"])
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                doc.add_paragraph(f"图片不存在: {image_path}")
        elif line.startswith("---"):
            doc.add_paragraph("─" * 40)
        elif line.strip():
            doc.add_paragraph(line.strip())
        i += 1
    doc.save(output_path)


# ==================== 3. 变化图斑矢量导出 ====================

@register_tool("report")
@tool
def export_change_vector(change_geojson_path: str, output_format: str = "shapefile", extract_shp: bool = False) -> Dict[str, Any]:
    """
    将变化图斑矢量数据导出为 Shapefile (ZIP) 或 GeoJSON, 供下游 GIS 系统使用 (#17)。

    ★ 实现: geopandas 读 GeoJSON → 写 Shapefile/GeoJSON, Shapefile 多文件打包 ZIP。
    ★ extract_shp=True: shapefile 模式下, 额外解压 ZIP 把 .shp 留在输出目录,
      返回 data.shp_path (绝对路径), 便于后续 upload_shapefile_layer 直接使用。

    入参:
        - change_geojson_path (str): 变化图斑 GeoJSON 路径 (detect_change 产物)
        - output_format (str): 输出格式, "shapefile" (默认, ZIP 打包) / "geojson"
    出参: {type: frontend_action, instruction: download, summary, data, verification}

    ★ 成功契约 (全部满足才算完成, 否则视为失败):
      1. type == "frontend_action" (而非 "error")
      2. verification.ok == True
      3. shapefile 模式: ZIP 内必须含 .shp/.dbf/.shx 三个核心文件 (非空目录条目)
      4. geojson 模式: geopandas 能读回且 feature_count > 0
      5. data.feature_count > 0 (源图斑非空)
      extract_shp=True 时还会额外满足: data.shp_path 指向真实存在的 .shp 文件
      任一不满足 → 必须按失败处理, 不可向用户汇报"已导出"
    示例:
        export_change_vector("E:/.../change.geojson", "shapefile")
        export_change_vector("E:/.../change.geojson", "shapefile", extract_shp=True)
    """
    gdf = _load_change_geojson(change_geojson_path)
    if gdf is None:
        return {"type": "error", "msg": f"无法读取变化图斑: {change_geojson_path}"}
    if gdf.empty:
        return {"type": "error", "msg": "变化图斑为空"}

    out_dir = _get_report_output_dir()
    # ★ v2.4 命名精简: 不再吃上游长 stem (套娃根因), 改用 {conv短}_export_{随机4}
    #   ASCII 化仍保留 (Fiona ESRI Shapefile 驱动中文 stem 会静默写出空文件)
    safe_stem = ascii_stem(short_conv() or "anon", fallback="anon") + "_export"
    rs = rand_suffix(4)
    output_format = output_format.lower()

    try:
        if output_format == "geojson":
            out_path = out_dir / f"{safe_stem}_{rs}.geojson"
            gdf.to_file(str(out_path), driver="GeoJSON", encoding="utf-8")
        elif output_format == "shapefile":
            # 先写出 shapefile 多文件, 再打包 ZIP
            shp_dir = out_dir / f"{safe_stem}_{rs}_shp"
            shp_dir.mkdir(parents=True, exist_ok=True)
            # ★ 修复 0KB bug 根因: Fiona 1.10+ 要求传带 .shp 扩展名的完整路径,
            #   无扩展名会被当成 GDAL dataset 目录 (创建空目录而非文件)。
            shp_base = shp_dir / f"{safe_stem}.shp"
            gdf.to_file(str(shp_base), driver="ESRI Shapefile", encoding="utf-8")
            # ★ 写出后立即校验 .shp 文件真实落地 (Fiona 静默失败时不会生成)
            if not shp_base.is_file() or shp_base.stat().st_size == 0:
                import shutil
                shutil.rmtree(str(shp_dir), ignore_errors=True)
                return {
                    "type": "error",
                    "msg": (f"Shapefile 写入失败: .shp 文件未生成或为空 "
                            f"(可能因字段名超长/几何无效/编码问题, stem={safe_stem})"),
                    "verification": {"ok": False, "reason": ".shp 文件未落地",
                                     "evidence": "Fiona 静默失败, ZIP 打包前已拦截"},
                }
            # ZIP 打包
            import zipfile
            zip_path = out_dir / f"{safe_stem}_{rs}.zip"
            with zipfile.ZipFile(str(zip_path), "w", zipfile.ZIP_DEFLATED) as zf:
                for f in shp_dir.iterdir():
                    if f.is_file():  # 只打包文件, 跳过子目录
                        zf.write(str(f), f.name)
            # 删除中间 shp 目录
            import shutil
            shutil.rmtree(str(shp_dir), ignore_errors=True)
            out_path = zip_path
        else:
            return {"type": "error", "msg": f"output_format 无效: {output_format} (shapefile/geojson)"}
    except Exception as e:
        return {"type": "error", "msg": f"矢量导出失败: {e}"}

    # ★ 成功契约校验: shapefile 验 ZIP 完整性, geojson 验要素数
    verification = verify_by_path(out_path)

    if not verification["ok"]:
        # ★ 不准伪装成功: 校验失败直接报错
        return {
            "type": "error",
            "msg": f"矢量导出失败: {verification.get('reason', '校验未通过')}",
            "verification": verification,
            "data": {"output_path": str(out_path), "format": output_format,
                     "feature_count": len(gdf)},
        }

    download_url = _build_download_url(out_path) + f"?t={rs}"
    size = verification.get("size_bytes", out_path.stat().st_size)

    # ★ extract_shp=True: 解压 ZIP, 把 .shp 留在输出目录, 供 upload_shapefile_layer 直接用
    shp_path: Optional[str] = None
    if extract_shp and output_format == "shapefile":
        try:
            import zipfile
            extract_dir = out_path.parent / f"{safe_stem}_{rs}_extracted"
            extract_dir.mkdir(parents=True, exist_ok=True)
            with zipfile.ZipFile(str(out_path), "r") as zf:
                zf.extractall(str(extract_dir))
            # 找 .shp 文件
            shp_candidates = list(extract_dir.glob("*.shp"))
            if shp_candidates:
                shp_path = str(shp_candidates[0].resolve())
                logger.info(f"[Export] extract_shp=True: .shp 解压到 {shp_path}")
            else:
                logger.warning("[Export] extract_shp=True 但 ZIP 内无 .shp")
        except Exception as e:
            logger.warning(f"[Export] extract_shp 解压失败: {e}")

    summary = (
        f"已导出变化图斑矢量: {out_path.name} ({output_format}, {size//1024}KB, {len(gdf)} 个要素)\n"
        f"下载链接已生成。"
        + (f"\n.shp 已解压: {shp_path} (可直接用于 upload_shapefile_layer)" if shp_path else "")
    )

    data = {
        "output_path": str(out_path),
        "output_url": download_url,
        "format": output_format,
        "feature_count": len(gdf),
    }
    if shp_path:
        data["shp_path"] = shp_path

    return build_download_action(
        params={"url": download_url, "filename": out_path.name, "caption": f"变化图斑 ({output_format})"},
        summary=summary,
        description=f"⏳ 矢量导出完成 ({output_format})",
        verification=verification,
        data=data,
    )


# ==================== 4. 统计表格导出 ====================

@register_tool("report")
@tool
def export_stats_table(change_geojson_path: str, output_format: str = "excel") -> Dict[str, Any]:
    """
    将变化图斑的统计数据导出为 Excel 或 CSV 表格 (#17)。

    ★ 表格字段: class_name / area_m2 / perimeter_m / 各属性列
    ★ 实现: pandas DataFrame → Excel (openpyxl) / CSV。

    入参:
        - change_geojson_path (str): 变化图斑 GeoJSON 路径
        - output_format (str): 输出格式, "excel" (默认) / "csv"
    出参: {type: frontend_action, instruction: download, summary, data, verification}

    ★ 成功契约 (全部满足才算完成, 否则视为失败):
      1. type == "frontend_action" (而非 "error")
      2. verification.ok == True (文件可被读回且非空)
      3. data.row_count > 0 (有数据行)
      任一不满足 → 必须按失败处理, 不可向用户汇报"已导出"
    示例:
        export_stats_table("E:/.../change.geojson", "excel")
    """
    gdf = _load_change_geojson(change_geojson_path)
    if gdf is None:
        return {"type": "error", "msg": f"无法读取变化图斑: {change_geojson_path}"}
    if gdf.empty:
        return {"type": "error", "msg": "变化图斑为空"}

    try:
        import pandas as pd
    except ImportError:
        return {"type": "error", "msg": "pandas 未安装"}

    # 提取属性列 (去掉 geometry)
    prop_cols = [c for c in gdf.columns if c != gdf.geometry.name]
    df = gdf[prop_cols].copy()
    # 加一列面积 (如果有 area_m2 用, 否则现算)
    if "area_m2" not in df.columns:
        try:
            df["area_m2"] = gdf.geometry.area.round(2)
        except Exception:
            pass
    # 加一列周长
    if "perimeter_m" not in df.columns:
        try:
            df["perimeter_m"] = gdf.geometry.length.round(2)
        except Exception:
            pass

    out_dir = _get_report_output_dir()
    # ★ v2.4 命名精简: {conv短}_stats_{随机4}
    safe_stem = ascii_stem(short_conv() or "anon", fallback="anon") + "_stats"
    rs = rand_suffix(4)
    output_format = output_format.lower()

    try:
        if output_format == "csv":
            out_path = out_dir / f"{safe_stem}_{rs}.csv"
            df.to_csv(str(out_path), index=False, encoding="utf-8-sig")
        elif output_format == "excel":
            out_path = out_dir / f"{safe_stem}_{rs}.xlsx"
            df.to_excel(str(out_path), index=False, sheet_name="变化统计")
        else:
            return {"type": "error", "msg": f"output_format 无效: {output_format} (excel/csv)"}
    except Exception as e:
        return {"type": "error", "msg": f"表格导出失败: {e}"}

    # ★ 成功契约校验
    verification = verify_by_path(out_path)
    if not verification["ok"]:
        return {
            "type": "error",
            "msg": f"表格导出失败: {verification.get('reason', '校验未通过')}",
            "verification": verification,
            "data": {"output_path": str(out_path), "format": output_format, "row_count": len(df)},
        }

    download_url = _build_download_url(out_path) + f"?t={rs}"
    size = verification.get("size_bytes", out_path.stat().st_size)

    summary = (
        f"已导出统计表格: {out_path.name} ({output_format}, {size//1024}KB, {len(df)} 行)\n"
        f"字段: {', '.join(df.columns.tolist())}\n"
        f"下载链接已生成。"
    )

    return build_download_action(
        params={"url": download_url, "filename": out_path.name, "caption": f"变化统计表 ({output_format})"},
        summary=summary,
        description=f"⏳ 表格导出完成 ({output_format})",
        verification=verification,
        data={
            "artifact_path": str(out_path),
            "artifact_url": download_url,
            "format": output_format,
            "row_count": len(df),
            "columns": df.columns.tolist(),
        },
    )
